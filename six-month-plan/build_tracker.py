#!/usr/bin/env python3
"""Build tracker.xlsx for the six-month consistency plan (1 Oct 2026 - 4 Apr 2027).

Usage:  python3 build_tracker.py            (requires: pip install openpyxl)

The workbook is formula driven: tick 1/0 in DailyLog and the weekly and monthly
completion percentages, status flags and dashboard update themselves.
"""
from __future__ import annotations

import math
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

START = date(2026, 10, 1)          # Day 1 (Thursday)
WEEK1_MONDAY = date(2026, 10, 5)   # Week 1 starts on a Monday so Mon/Wed/Fri gym days line up
END = date(2027, 4, 4)             # Sunday, end of Week 26
N_DAYS = (END - START).days + 1    # 186
N_WEEKS = 26
DL_FIRST, DL_LAST = 2, N_DAYS + 1  # DailyLog data rows

OUT = Path(__file__).with_name("tracker.xlsx")

# ---------------------------------------------------------------- styles
NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
SUNDAY_FILL = PatternFill("solid", fgColor="EDEDED")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
BOLD = Font(bold=True)
DATE_FMT = "ddd dd-mmm-yy"


def col_width(ws, idx: int) -> float:
    w = ws.column_dimensions[get_column_letter(idx)].width
    return w if w else 8.43


def autoheight(ws, row: int, first_col: int, last_col: int, spans: dict | None = None):
    """Approximate row height from the longest wrapped text in the row."""
    lines = 1
    for c in range(first_col, last_col + 1):
        v = ws.cell(row=row, column=c).value
        if v is None:
            continue
        end = (spans or {}).get(c, c)
        width = sum(col_width(ws, k) for k in range(c, end + 1))
        chars = max(1, int(width * 1.15))
        n = sum(max(1, math.ceil(len(part) / chars)) for part in str(v).split("\n"))
        lines = max(lines, n)
    ws.row_dimensions[row].height = min(409, 15 * lines + 3)


def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def set_widths(ws, widths: dict[str, float]):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def write_table(ws, top, headers, rows, start_col=1, merge_to=None, wrap=True):
    """Write a header row plus data rows at (top, start_col); return the next free row.

    merge_to: absolute column index; the last table column is merged out to it on every row.
    """
    last_col = start_col + len(headers) - 1
    spans = {last_col: merge_to} if merge_to else None
    for j, h in enumerate(headers):
        ws.cell(row=top, column=start_col + j, value=h)
    style_header(ws, top, start_col, merge_to or last_col)
    if merge_to:
        ws.merge_cells(start_row=top, start_column=last_col, end_row=top, end_column=merge_to)
    r = top + 1
    for row in rows:
        for j, v in enumerate(row):
            cell = ws.cell(row=r, column=start_col + j, value=v)
            cell.border = BORDER
            if wrap:
                cell.alignment = WRAP
        if merge_to:
            for k in range(last_col + 1, merge_to + 1):
                ws.cell(row=r, column=k).border = BORDER
            ws.merge_cells(start_row=r, start_column=last_col, end_row=r, end_column=merge_to)
        if wrap:
            autoheight(ws, r, start_col, last_col, spans)
        r += 1
    return r + 1


def paragraph(ws, row, text, first_col=1, last_col=6, bold=False):
    c = ws.cell(row=row, column=first_col, value=text)
    c.alignment = WRAP
    if bold:
        c.font = BOLD
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    autoheight(ws, row, first_col, first_col, {first_col: last_col})


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="595959")


# ---------------------------------------------------------------- calendar logic
def week_of(d: date) -> int:
    if d < WEEK1_MONDAY:
        return 0
    return (d - WEEK1_MONDAY).days // 7 + 1


def cycle_of(week: int) -> int:
    return 1 if week <= 13 else 2


def phase_of(week: int) -> str:
    if week == 0:
        return "Launch"
    if week <= 2:
        return "Foundation (technique)"
    if week <= 8:
        return "Build 1"
    if week == 9:
        return "Deload"
    if week <= 12:
        return "Build 2 (holiday stress test)"
    if week == 13:
        return "Benchmarks + Cycle 1 review"
    if week <= 17:
        return "Cycle 2 block A"
    if week <= 21:
        return "Cycle 2 block B"
    if week == 22:
        return "Deload + diet break"
    if week <= 25:
        return "Peak"
    return "Final benchmarks + review"


NUTRITION_BY_PHASE = {
    "Launch": "Set up kitchen, first meal prep",
    "Foundation (technique)": "Deficit starts: 2,000-2,100 kcal, 140 g protein",
    "Build 1": "Same; 2 planned sweets per week from November",
    "Deload": "Same",
    "Build 2 (holiday stress test)": "Hold the minimums through the holidays",
    "Benchmarks + Cycle 1 review": "Checkpoint: waist < 87 cm -> maintenance, else continue deficit",
    "Cycle 2 block A": "Per checkpoint",
    "Cycle 2 block B": "Per checkpoint",
    "Deload + diet break": "Maintenance for 7 days",
    "Peak": "Finish deficit or maintain",
    "Final benchmarks + review": "Decide the next 6 months",
}


def week_dates(week: int) -> tuple[date, date]:
    if week == 0:
        return START, WEEK1_MONDAY - timedelta(days=1)
    s = WEEK1_MONDAY + timedelta(days=7 * (week - 1))
    return s, s + timedelta(days=6)


def planned_training(d: date) -> str:
    wd = d.weekday()  # Monday = 0
    week = week_of(d)
    if week == 0:
        return {
            3: "Trainer intake + baseline tests (push-ups, dead hang, plank, goblet squat 10RM, resting HR), photos, measurements",
            4: "Strength A (technique loads)",
            5: "Activity 45-60 min",
            6: "Mobility 20 min + finish baseline measurements",
        }[wd]
    if week <= 13:
        if wd in (0, 2, 4):
            if week == 13:
                return {
                    0: "Benchmarks 1: back squat 5RM, bench press 5RM, push-ups to failure",
                    2: "Benchmarks 2: deadlift 5RM, strict chin-ups, dead hang, plank",
                    4: "Light full body + 12-min Cooper test",
                }[wd]
            idx = {0: 0, 2: 1, 4: 2}[wd]
            letter = "A" if (3 * (week - 1) + idx) % 2 == 0 else "B"
            s = f"Strength {letter}"
            if week <= 2:
                s += " (2 sets, RPE 5-6)"
            elif week == 9:
                s += " (deload: 2 sets, -40% load)"
            elif week <= 8:
                s += " (3 sets, RPE 7-8)"
            else:
                s += " (3-4 sets, main lifts 6-10 reps)"
            return s
        if wd in (1, 3):
            return "Zone 2 cardio 30 min" if week <= 4 else "Zone 2 cardio 35-40 min"
        if wd == 5:
            return "Activity 45-60 min (hike, sport, long walk)"
        if week == 13:
            return "Cycle 1 review + measurements + photos + rewrite rules"
        return "Mobility 20 min + weekly review"
    names = {0: "Lower A", 1: "Upper A", 2: "Zone 2 cardio 40 min", 3: "Lower B", 4: "Upper B",
             5: "Intervals 8 x 1 min hard / 2 min easy, or sport 45-60 min", 6: "Mobility 20 min + weekly review"}
    if week == 26:
        return {
            0: "Final benchmarks 1: back squat 5RM, bench press 5RM, push-ups to failure",
            1: "Upper A (light)",
            2: "Zone 2 + 12-min Cooper test",
            3: "Final benchmarks 2: deadlift 5RM, strict chin-ups, dead hang, plank",
            4: "Upper B (light)",
            5: "Activity 45-60 min",
            6: "Final review + measurements + photos + next 6-month plan",
        }[wd]
    s = names[wd]
    if wd in (0, 1, 3, 4):
        if week == 22:
            s += " (deload: 2 sets, -40% load)"
        elif week >= 23:
            s += " (peak: main lifts 4 x 5-6)"
    return s


SKILL_PLAN = [
    # week, module, focus, resources, friday deliverable
    (0, "Setup", "Install Miniforge/conda env (geopandas, rasterio, rioxarray, xarray, rasterstats, matplotlib, jupyterlab, scikit-learn, earthengine-api, geemap), VS Code, Git. Create the public portfolio repo. Register for Google Earth Engine (approval can take days). Choose the study area and download admin boundaries, roads, DEM.", "Miniforge docs; GitHub; earthengine.google.com", "Portfolio repo with README: study area + 26-week plan"),
    (1, "Python for geospatial", "Python and pandas refresh: types, functions, DataFrames, reading CSV, plotting.", "Spatial Thoughts - Python Foundation for Spatial Analysis (part 1)", "Notebook 01: pandas analysis of an attribute table exported from ArcGIS Pro"),
    (2, "Python for geospatial", "GeoPandas vectors: read shapefile/GDB, CRS and reprojection (EPSG:4326 vs UTM), buffer, overlay, spatial join, dissolve.", "Spatial Thoughts - Python Foundation (GeoPandas); GeoPandas docs", "Notebook 02: vector analysis of the study area (buffers + spatial join) with a map"),
    (3, "Python for geospatial", "Rasters with rasterio/rioxarray: open DEM, reproject, clip to boundary, slope and hillshade with numpy, zonal statistics (rasterstats).", "Rasterio and rioxarray docs; rasterstats docs", "Notebook 03: terrain statistics per district"),
    (4, "Python for geospatial", "Integration project: reproduce one ArcGIS Pro Day-1 exercise end to end (buffers, overlay, zonal stats), export to GeoPackage, write README.", "Own course material", "Project 1 folder: README, figures, GeoPackage"),
    (5, "Google Earth Engine", "Code Editor, Image and ImageCollection, filters, Sentinel-2 cloud masking (SCL / Cloud Score+), median composites, visualisation.", "Spatial Thoughts - End-to-End GEE (modules 1-2); EEFA open-access book F1", "Script 01: cloud-free Sentinel-2 composite of the study area"),
    (6, "Google Earth Engine", "Reducers and indices: NDVI/NDWI/NDBI, zonal statistics by admin unit (reduceRegions), exports to Drive/Assets.", "End-to-End GEE (modules 2-3); EEFA F2-F3", "Script 02: mean NDVI per district table + map export"),
    (7, "Google Earth Engine", "Time series: monthly NDVI 2019-2026, charts (ui.Chart), simple anomalies, GEE Apps basics.", "End-to-End GEE (modules 3-4); EEFA F4", "NDVI time-series chart + a published GEE App"),
    (8, "Google Earth Engine", "GEE Python API and geemap: authenticate, replicate the scripts in Jupyter, interactive maps, export GeoTIFF.", "End-to-End GEE (Python module); geemap docs", "Notebook 04: GEE composites and time series in Python"),
    (9, "RS analytics (deload week: theory)", "Supervised vs unsupervised classification, spectral signatures, training-sample design, Random Forest, accuracy metrics.", "EEFA F2.1-F2.2; NASA ARSET - Using GEE for Land Monitoring (part 1)", "One-page methods note + training sample scheme (5-6 classes)"),
    (10, "RS analytics", "Random Forest classification in GEE: collect training samples, train, classify the Sentinel-2 composite, visualise.", "EEFA F2.1; End-to-End GEE supervised classification", "Land-cover map v1 of the study area"),
    (11, "RS analytics", "Accuracy assessment: stratified random validation points, confusion matrix, overall/producer/user accuracy, F1; add features (indices, texture, DEM) and iterate.", "EEFA F2.2; ARSET GEE part 2 (classification and accuracy)", "Accuracy report + land-cover map v2 (target >= 85% OA)"),
    (12, "RS analytics", "Change detection: classify a 2016 composite, post-classification comparison, change matrix, area statistics per district; repeat in Python with scikit-learn on an exported stack.", "EEFA F4.4 / A1; scikit-learn RandomForest docs", "Change map 2016-2026 + change matrix table"),
    (13, "Capstone 1", "Write-up: methods, maps, tables, limitations in the README; publish; review the skill plan for Cycle 2.", "Own results", "Capstone 1 published in the portfolio"),
    (14, "ArcPy and automation", "ArcPy basics: env settings, Describe, arcpy.da Search/Update cursors, running geoprocessing tools from scripts, logging.", "Esri ArcPy docs; Esri Academy free Python web courses", "batch_project_clip.py over a folder of feature classes"),
    (15, "ArcPy and automation", "Geodatabase automation: create domains, subtypes and relationship classes in code; QA checks (empty geometry, missing attributes, CRS) with a report.", "Esri ArcPy data management docs", "gdb_qa_report.py producing a CSV/Excel report"),
    (16, "ArcPy and automation", "Python toolbox (.pyt): parameters, validation, messages; package for colleagues.", "Esri Python toolbox docs", "Toolbox v1 (QA report or coverage buffers from tower points) + usage README"),
    (17, "PostGIS and spatial SQL", "Install PostgreSQL + PostGIS, create DB, load data (ogr2ogr / QGIS DB Manager), spatial indexes, basic SELECTs.", "postgis.net - Introduction to PostGIS workshop (sections 1-9)", "Database loaded with study-area layers + 3 queries"),
    (18, "PostGIS and spatial SQL", "Spatial SQL: ST_Buffer, ST_Intersects, ST_Within, ST_DWithin, spatial joins, GROUP BY aggregation, KNN nearest neighbour.", "PostGIS workshop (sections 10-20)", "8 analytical queries with results in the README"),
    (19, "PostGIS and spatial SQL", "Integration: connect QGIS and ArcGIS Pro (query layers) to PostGIS, views, raster basics; compare with enterprise geodatabase concepts.", "QGIS / ArcGIS Pro docs", "Project 3 README + views + screenshots"),
    (20, "Web GIS and dashboards", "Publish classification and change results as hosted layers/tiles (ArcGIS Online) or COG + tiles. Book the certification exam and buy the voucher.", "ArcGIS Online docs; Esri Academy exam page", "Published layers + exam date booked"),
    (21, "Web GIS and dashboards", "Dashboard / web app: ArcGIS Dashboards or Experience Builder, or Streamlit + leafmap; filter by district, change-statistics charts.", "Spatial Thoughts - Mapping and Data Visualization with Python", "Public web app link"),
    (22, "Web GIS (deload week: polish)", "Polish repos and write a 600-900 word article about Capstone 1 (LinkedIn or blog).", "Own results", "Article published"),
    (23, "Capstone 2", "Pipeline: GEE classification (2016, 2026) -> export -> PostGIS -> toolbox automation.", "All previous modules", "Pipeline running end to end on the study area"),
    (24, "Capstone 2", "Validation (>= 85% OA), change statistics, web app updated, documentation.", "All previous modules", "Capstone 2 published"),
    (25, "Certification prep", "Esri exam learning plan, Exam Information Guide sections, practice questions, weak-area review.", "Esri Academy learning plan; Esri Community EIG", "Practice test >= 80%"),
    (26, "Exam + review", "Sit the exam; portfolio review; write the next six-month skill plan.", "-", "Exam taken; next plan committed"),
]
SKILL_MODULE = {w: m for (w, m, *_rest) in SKILL_PLAN}

RULES = [
    ("R1 Move", "Do the planned session (strength, Zone 2, Saturday activity or Sunday mobility) on today's row", "15-minute walk"),
    ("R2 Eat", "Protein at every meal, no sugary drinks or packaged snacks, 2.5-3 L water", "Protein at 2 meals and the water"),
    ("R3 Read", "10 pages of non-fiction", "5 pages"),
    ("R4 Mind", "5-minute journal: today's win, today's lesson, tomorrow's #1 task", "One line"),
    ("R5 Sleep", "Screens off 22:15, lights out 23:00, wake 06:30 (+/- 30 min)", "Phone charges outside the bedroom"),
    ("R6 Skill (Mon-Fri)", "60 minutes on the current GIS/RS module, timer on, phone away", "20 minutes"),
]


# ---------------------------------------------------------------- sheets
def build_overview(wb):
    ws = wb.active
    ws.title = "Overview"
    set_widths(ws, {"A": 22, "B": 70, "C": 40, "D": 48})
    title(ws, "Six-Month Consistency Plan - Tracker",
          f"Day 1: {START:%A %d %b %Y}   |   Week 1 starts {WEEK1_MONDAY:%d %b}   |   Last day: {END:%A %d %b %Y}   |   {N_DAYS} days = 4 launch days + 26 weeks")
    ws["A4"] = "Direction statement (who I am on 4 Apr 2027) - write 3 sentences:"
    ws["A4"].font = BOLD
    for r in (5, 6, 7):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(row=r, column=1).fill = INPUT_FILL
        ws.cell(row=r, column=1).border = BORDER
        ws.row_dimensions[r].height = 22
    ws["A9"] = "Accountability person:"
    ws["A9"].font = BOLD
    ws["B9"].fill = INPUT_FILL
    ws["A10"] = "Study area for the skill block:"
    ws["A10"].font = BOLD
    ws["B10"].fill = INPUT_FILL

    r = write_table(ws, 12, ["Rule", "Counts as done when", "Minimum version (still counts)"], RULES)
    paragraph(ws, r, "Scoring: 6 rules on weekdays, 5 on weekends. Target 80-90% per week. Below 70% two weeks running: use the minimum versions as the standard for a week. Above 95% three weeks running: raise the bar. A missed day stays missed and the challenge never restarts. Never miss twice in a row. Weekly rule: one commit pushed to the portfolio repo every Friday.", 1, 4)
    r += 2

    spans = []
    for w in range(0, N_WEEKS + 1):
        p = phase_of(w)
        if spans and spans[-1][0] == p:
            spans[-1][2] = w
        else:
            spans.append([p, w, w])
    phases = []
    for p, w1, w2 in spans:
        s, _ = week_dates(w1)
        _, e = week_dates(w2)
        label = f"{w1}" if w1 == w2 else f"{w1}-{w2}"
        phases.append((label, f"{s:%d %b} - {e:%d %b %Y}", p, NUTRITION_BY_PHASE[p]))
    r = write_table(ws, r, ["Weeks", "Dates", "Phase", "Nutrition"], phases)

    targets = [
        ("Waist at navel", "measure Day 1", "-5 cm", "-8 to -10 cm and below 87 cm"),
        ("Body weight", "74 kg", "70-71 kg", "67-70 kg"),
        ("Push-ups to failure", "measure Day 1", "+8", "+15 (minimum 25)"),
        ("Strict chin-ups", "measure Day 1", "first strict rep", "5"),
        ("Back squat 5RM", "learn W1-4", "0.75 x BW", "~1.0 x BW (~70 kg)"),
        ("Deadlift 5RM", "learn W1-4", "1.0 x BW", "~1.25 x BW (~90 kg)"),
        ("Bench press 5RM", "dumbbells first", "0.6 x BW", "~0.75 x BW (~52 kg)"),
        ("Dead hang", "measure Day 1", "45-60 s", "60-90 s"),
        ("Resting heart rate", "measure Day 1", "-3 bpm", "-5 to -8 bpm"),
    ]
    r = write_table(ws, r, ["Target", "Now", "Week 13 (3 Jan 2027)", "Week 26 (4 Apr 2027)"], targets)

    ws.cell(row=r, column=1, value="How to use this workbook").font = BOLD
    tips = [
        "DailyLog: every evening enter 1 or 0 for each rule (R6 only Mon-Fri). Steps, sleep hours and weight are optional. Everything else is computed.",
        "WeeklyReview: on Sunday add weight, waist, energy, the biggest friction and one adjustment. Completion % and status are computed from DailyLog.",
        "Benchmarks: fill the test columns on Day 1-4, then every 4 weeks; the full test battery in weeks 13 and 26.",
        "TrainingPlan / TrainingLog: give TrainingPlan to your trainer; log the key lifts per session in TrainingLog.",
        "SkillPlan: one row per week with the module, focus, resources and the Friday deliverable; mark Done and log hours.",
        "Targets: edit sex, age, weight and activity factor; calorie, protein and water targets recalculate. Meals references those cells.",
        "MonthlyReview: six rows of prompts; the month completion % is computed.",
        "Regenerate a blank copy any time with: python3 build_tracker.py",
    ]
    for i, t in enumerate(tips, start=1):
        paragraph(ws, r + i, f"{i}. {t}", 1, 4)


def build_targets(wb):
    ws = wb.create_sheet("Targets")
    set_widths(ws, {"A": 30, "B": 12, "C": 90})
    title(ws, "Targets calculator", "Yellow cells are inputs. Mifflin-St Jeor BMR x activity factor, minus the deficit; protein and fat per kg bodyweight.")
    inputs = [
        ("sex", "Sex (M/F)", "M", "Used by the BMR formula"),
        ("age", "Age (years)", 35, "Edit"),
        ("height", "Height (cm)", 174, ""),
        ("weight", "Current weight (kg)", 74, "Update monthly; targets recalculate"),
        ("af", "Activity factor", 1.45, "1.2 sedentary, 1.375 light, 1.45 desk job + 4-5 sessions/week + 10k steps, 1.55 moderate"),
        ("deficit", "Deficit (fraction)", 0.15, "0.15 = mild deficit (~300-400 kcal). Use 0 for maintenance (week 22 diet break, or once the waist target is reached)"),
        ("ppk", "Protein (g per kg)", 1.8, "1.6-2.2 g/kg while in a deficit"),
        ("fpk", "Fat (g per kg)", 0.8, "0.7-1.0 g/kg"),
    ]
    for j, h in enumerate(["Input", "Value", "Note"], start=1):
        ws.cell(row=3, column=j, value=h).font = BOLD
    ref = {}
    r = 4
    for key, label, val, note in inputs:
        ws.cell(row=r, column=1, value=label).border = BORDER
        c = ws.cell(row=r, column=2, value=val)
        c.fill = INPUT_FILL
        c.border = BORDER
        ws.cell(row=r, column=3, value=note).alignment = WRAP
        ref[key] = f"B{r}"
        r += 1
    dv = DataValidation(type="list", formula1='"M,F"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ref["sex"])

    r += 1
    for j, h in enumerate(["Output", "Value", "Note"], start=1):
        ws.cell(row=r, column=j, value=h).font = BOLD
    outputs = [
        ("bmi", "BMI", "={weight}/({height}/100)^2", "0.0", "18.5-24.9 normal; yours is normal, the issue is fat distribution"),
        ("bmr", "BMR (kcal)", '=IF({sex}="M",10*{weight}+6.25*{height}-5*{age}+5,10*{weight}+6.25*{height}-5*{age}-161)', "0", "Mifflin-St Jeor"),
        ("tdee", "Maintenance / TDEE (kcal)", "={bmr}*{af}", "0", ""),
        ("kcal", "Calorie target (kcal)", "={tdee}*(1-{deficit})", "0", "Eat within +/- 100 kcal of this on average"),
        ("protein", "Protein target (g)", "={ppk}*{weight}", "0", "Spread over 4-5 meals"),
        ("fat", "Fat target (g)", "={fpk}*{weight}", "0", ""),
        ("carbs", "Carbohydrate target (g)", "=({kcal}-4*{protein}-9*{fat})/4", "0", "Mostly at breakfast, pre-gym and dinner"),
        ("water_rest", "Water target, rest day (L)", "=0.035*{weight}", "0.0", "35 ml per kg"),
        ("water", "Water target, training day (L)", "=0.035*{weight}+0.5", "0.0", "Rest-day amount + 0.5 L around the session"),
        ("waist", "Waist target (cm)", "=0.5*{height}", "0", "Waist-to-height ratio below 0.5"),
        ("loss", "Expected loss per week (kg)", "={tdee}*{deficit}*7/7700", "0.00", "Scale may move slower if muscle is gained; trust the waist"),
    ]
    for key, label, formula, fmt, note in outputs:
        r += 1
        ref[key] = f"B{r}"
        ws.cell(row=r, column=1, value=label).border = BORDER
        c = ws.cell(row=r, column=2, value=formula.format(**ref))
        c.number_format = fmt
        c.border = BORDER
        c.font = BOLD
        ws.cell(row=r, column=3, value=note).alignment = WRAP
    return ref


def build_daily_log(wb):
    ws = wb.create_sheet("DailyLog")
    headers = ["Day", "Date", "Weekday", "Week", "Cycle", "Phase", "Planned training", "Skill module",
               "R1 Move", "R2 Eat", "R3 Read", "R4 Mind", "R5 Sleep", "R6 Skill",
               "Done", "Required", "Logged", "Req if logged", "Daily %",
               "Steps", "Sleep (h)", "Weight (kg)", "Notes"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, 1, len(headers))
    ws.row_dimensions[1].height = 32
    for i in range(N_DAYS):
        d = START + timedelta(days=i)
        r = DL_FIRST + i
        w = week_of(d)
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=d).number_format = DATE_FMT
        ws.cell(row=r, column=3, value=d.strftime("%a"))
        ws.cell(row=r, column=4, value=w)
        ws.cell(row=r, column=5, value=cycle_of(w))
        ws.cell(row=r, column=6, value=phase_of(w))
        ws.cell(row=r, column=7, value=planned_training(d))
        ws.cell(row=r, column=8, value=SKILL_MODULE[w])
        ws.cell(row=r, column=15, value=f"=SUM(I{r}:M{r})+IF(WEEKDAY(B{r},2)<=5,N{r},0)")
        ws.cell(row=r, column=16, value=f"=IF(WEEKDAY(B{r},2)<=5,6,5)")
        ws.cell(row=r, column=17, value=f"=IF(COUNT(I{r}:N{r})>0,1,0)")
        ws.cell(row=r, column=18, value=f"=P{r}*Q{r}")
        ws.cell(row=r, column=19, value=f'=IF(Q{r}=1,O{r}/P{r},"")').number_format = "0%"
        for col in range(9, 15):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        for col in (20, 21, 22, 23):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        if d.weekday() >= 5:
            ws.cell(row=r, column=14).fill = GREY_FILL
        if d.weekday() == 6:
            for col in (1, 2, 3):
                ws.cell(row=r, column=col).fill = SUNDAY_FILL
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
    rng = f"I{DL_FIRST}:N{DL_LAST}"
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["1"], fill=GREEN_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0"], fill=RED_FILL))
    pct = f"S{DL_FIRST}:S{DL_LAST}"
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(S{DL_FIRST}),S{DL_FIRST}>=0.8)"], fill=GREEN_FILL))
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(S{DL_FIRST}),S{DL_FIRST}<0.8,S{DL_FIRST}>=0.6)"], fill=AMBER_FILL))
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(S{DL_FIRST}),S{DL_FIRST}<0.6)"], fill=RED_FILL))
    widths = [5, 14, 8, 6, 6, 26, 46, 30, 8, 8, 8, 8, 8, 8, 6, 8, 7, 8, 8, 8, 9, 10, 40]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for col in ("Q", "R"):
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "D2"


def build_weekly_review(wb):
    ws = wb.create_sheet("WeeklyReview")
    title(ws, "Weekly review (Sundays, 15 minutes)")
    hr = 7
    first, last = hr + 1, hr + 1 + N_WEEKS
    dash = [
        ("Days logged", f"=SUM(DailyLog!Q{DL_FIRST}:Q{DL_LAST})", "0"),
        ("Overall completion to date", f'=IFERROR(SUM(DailyLog!O{DL_FIRST}:O{DL_LAST})/SUM(DailyLog!R{DL_FIRST}:R{DL_LAST}),"")', "0%"),
        ("Current week (by today's date)", f"=IF(TODAY()<DATE({WEEK1_MONDAY.year},{WEEK1_MONDAY.month},{WEEK1_MONDAY.day}),0,MIN({N_WEEKS},INT((TODAY()-DATE({WEEK1_MONDAY.year},{WEEK1_MONDAY.month},{WEEK1_MONDAY.day}))/7)+1))", "0"),
        ("Weeks at or above 80%", f'=COUNTIF(H{first}:H{last},">=0.8")', "0"),
    ]
    for i, (label, f, fmt) in enumerate(dash, start=2):
        ws.cell(row=i, column=1, value=label).font = BOLD
        c = ws.cell(row=i, column=2, value=f)
        c.number_format = fmt
        c.font = BOLD
        c.fill = SUB_FILL
    headers = ["Week", "Start", "End", "Cycle", "Phase", "Done", "Required (logged days)", "Completion %", "Status",
               "Training sessions (R1 ticks)", "Skill days (R6 ticks)", "Avg sleep (h)", "Avg steps",
               "Weight (kg)", "Waist (cm)", "Friday commit pushed (Y/N)", "Energy (1-5)", "Biggest friction", "Adjustment for next week"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=j, value=h)
    style_header(ws, hr, 1, len(headers))
    ws.row_dimensions[hr].height = 45
    dl = f"DailyLog!$D${DL_FIRST}:$D${DL_LAST}"
    for w in range(0, N_WEEKS + 1):
        r = hr + 1 + w
        s, e = week_dates(w)
        ws.cell(row=r, column=1, value=w)
        ws.cell(row=r, column=2, value=s).number_format = DATE_FMT
        ws.cell(row=r, column=3, value=e).number_format = DATE_FMT
        ws.cell(row=r, column=4, value=cycle_of(w))
        ws.cell(row=r, column=5, value=phase_of(w))
        ws.cell(row=r, column=6, value=f"=SUMIFS(DailyLog!$O${DL_FIRST}:$O${DL_LAST},{dl},$A{r})")
        ws.cell(row=r, column=7, value=f"=SUMIFS(DailyLog!$R${DL_FIRST}:$R${DL_LAST},{dl},$A{r})")
        ws.cell(row=r, column=8, value=f'=IF(G{r}=0,"",F{r}/G{r})').number_format = "0%"
        ws.cell(row=r, column=9, value=f'=IF(H{r}="","",IF(H{r}>=0.95,"Raise the bar",IF(H{r}>=0.8,"On track",IF(H{r}>=0.7,"Watch","Shrink to minimums"))))')
        ws.cell(row=r, column=10, value=f"=COUNTIFS({dl},$A{r},DailyLog!$I${DL_FIRST}:$I${DL_LAST},1)")
        ws.cell(row=r, column=11, value=f"=COUNTIFS({dl},$A{r},DailyLog!$N${DL_FIRST}:$N${DL_LAST},1)")
        ws.cell(row=r, column=12, value=f'=IFERROR(AVERAGEIFS(DailyLog!$U${DL_FIRST}:$U${DL_LAST},{dl},$A{r}),"")').number_format = "0.0"
        ws.cell(row=r, column=13, value=f'=IFERROR(AVERAGEIFS(DailyLog!$T${DL_FIRST}:$T${DL_LAST},{dl},$A{r}),"")').number_format = "#,##0"
        for col in range(14, 20):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
        ws.cell(row=r, column=18).alignment = WRAP
        ws.cell(row=r, column=19).alignment = WRAP
    pct = f"H{first}:H{last}"
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(H{first}),H{first}>=0.8)"], fill=GREEN_FILL))
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(H{first}),H{first}<0.8,H{first}>=0.7)"], fill=AMBER_FILL))
    ws.conditional_formatting.add(pct, FormulaRule(formula=[f"AND(ISNUMBER(H{first}),H{first}<0.7)"], fill=RED_FILL))
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_energy = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    ws.add_data_validation(dv_yn)
    ws.add_data_validation(dv_energy)
    dv_yn.add(f"P{first}:P{last}")
    dv_energy.add(f"Q{first}:Q{last}")
    widths = [6, 14, 14, 6, 28, 7, 10, 11, 18, 10, 9, 9, 9, 9, 9, 10, 8, 30, 34]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"B{first}"


def build_benchmarks(wb):
    ws = wb.create_sheet("Benchmarks")
    set_widths(ws, {"A": 36, "B": 7, "C": 16, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13, "J": 20, "K": 20})
    title(ws, "Benchmarks and measurements", "Same conditions each time: morning, before breakfast, waist relaxed at the navel. Photos front/side/back in the same light and clothing.")
    headers = ["Metric", "Unit", "Day 1-4 (1-4 Oct)", "Wk 4 (1 Nov)", "Wk 8 (29 Nov)", "Wk 13 (3 Jan)",
               "Wk 17 (31 Jan)", "Wk 21 (28 Feb)", "Wk 26 (4 Apr)", "Target Wk 26", "Change (Wk 26 - Day 1)"]
    metrics = [
        ("Body weight", "kg", "67-70", True),
        ("Waist at navel", "cm", "-8 to -10 and < 87", True),
        ("Chest", "cm", "-", True),
        ("Hips", "cm", "-", True),
        ("Resting heart rate", "bpm", "-5 to -8", True),
        ("Push-ups to failure", "reps", "+15 (min 25)", True),
        ("Dead hang", "s", "60-90", True),
        ("Plank", "s", "90+", True),
        ("Strict chin-ups", "reps", "5", True),
        ("Goblet squat 10RM (Day 1 only)", "kg", "-", True),
        ("Back squat 5RM", "kg", "~70 (1.0 x BW)", True),
        ("Deadlift 5RM", "kg", "~90 (1.25 x BW)", True),
        ("Bench press 5RM", "kg", "~52 (0.75 x BW)", True),
        ("12-min Cooper test distance", "m", "+300 vs Day 1", True),
        ("Average daily steps (last 2 weeks)", "steps", "10,000", True),
        ("Average sleep (last 2 weeks)", "h", "7.5", True),
        ("Books finished (cumulative)", "count", "6", True),
        ("Skill deliverables shipped (cumulative)", "count", "26", True),
        ("Photos taken (Y/N)", "Y/N", "Y", False),
    ]
    hr = 4
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=j, value=h)
    style_header(ws, hr, 1, len(headers))
    ws.row_dimensions[hr].height = 34
    for i, (m, unit, target, numeric) in enumerate(metrics, start=1):
        r = hr + i
        ws.cell(row=r, column=1, value=m)
        ws.cell(row=r, column=2, value=unit)
        for col in range(3, 10):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.cell(row=r, column=10, value=target)
        if numeric:
            ws.cell(row=r, column=11, value=f'=IF(OR(C{r}="",I{r}=""),"",I{r}-C{r})')
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
    ws.freeze_panes = f"C{hr + 1}"


def build_training_plan(wb):
    ws = wb.create_sheet("TrainingPlan")
    set_widths(ws, {"A": 22, "B": 34, "C": 46, "D": 14, "E": 44})
    title(ws, "Training programme (proposal for the trainer)", "Structure, progression and grip strategy are what matter; the trainer chooses exact exercises and loads.")
    r = 4
    paragraph(ws, r, "Cycle 1 (Weeks 1-13): three full-body sessions Mon/Wed/Fri, alternating A/B (W1 = A,B,A; W2 = B,A,B ...). Warm-up 8 min: bike/rower, wrist circles + light wrist flexion/extension, band pull-aparts, 10 bodyweight squats.", 1, 5, bold=True)
    r += 2
    ws.cell(row=r, column=1, value="Session A").font = BOLD
    r = write_table(ws, r + 1, ["Slot", "Weeks 1-4", "Weeks 5-13", "Sets x reps", "Notes"], [
        ("Squat pattern", "Goblet squat", "Barbell back squat (or leg press)", "3 x 8-12", "Main lower-body driver"),
        ("Horizontal push", "Dumbbell bench press, neutral grip", "Barbell bench from week 9 if wrists are comfortable; stacked wrists", "3 x 8-12", "Neutral grip protects small wrists early"),
        ("Hip hinge", "Dumbbell Romanian deadlift", "Barbell RDL or trap-bar deadlift; straps allowed from week 8", "3 x 8-12", "Grip must never limit the hinge"),
        ("Horizontal pull", "Chest-supported dumbbell row", "Seated cable row", "3 x 10-12", "Posture"),
        ("Core", "Plank", "Long-lever or weighted plank", "3 x 30-60 s", ""),
        ("Grip finisher", "Farmer carry", "Farmer carry, heavier", "3 x 30-40 m", "Forearm size is what changes the look of a small wrist"),
    ])
    ws.cell(row=r, column=1, value="Session B").font = BOLD
    r = write_table(ws, r + 1, ["Slot", "Weeks 1-4", "Weeks 5-13", "Sets x reps", "Notes"], [
        ("Single-leg", "Split squat or walking lunge", "Bulgarian split squat", "3 x 8-10 / leg", ""),
        ("Vertical pull", "Lat pulldown", "Assisted chin-up, aim at the first strict rep by week 13", "3 x 8-12", ""),
        ("Vertical push", "Seated dumbbell shoulder press, neutral grip", "Same, heavier", "3 x 8-12", ""),
        ("Glute / hamstring", "Glute bridge", "Hip thrust or 45-degree back extension", "3 x 10-12", ""),
        ("Upper back / posture", "Face pull", "Face pull or rear-delt fly", "3 x 12-15", "Counteracts desk posture"),
        ("Core", "Side plank", "Pallof press", "3 x 30 s / 10 reps", ""),
        ("Grip finisher", "Dead hang", "Dead hang, target 60 s by week 13", "3 x max time", ""),
    ])
    ws.cell(row=r, column=1, value="Loading rules, cardio, steps").font = BOLD
    r = write_table(ws, r + 1, ["Weeks", "Rule"], [
        ("1-2", "2 sets per exercise, RPE 5-6 (4-5 reps in reserve). Technique and habit."),
        ("3-8", "3 sets, RPE 7-8. Double progression: top of the rep range on all sets -> +2.5 kg upper / +5 kg lower, restart at the bottom of the range."),
        ("9", "Deload: 2 sets, 40% less load, same movements."),
        ("10-12", "3-4 sets; main lifts 6-10 reps at RPE 8; accessories 10-15."),
        ("13", "Benchmarks. Mon: back squat 5RM, bench 5RM, push-ups. Wed: deadlift 5RM, chin-ups, dead hang, plank. Fri: light full body + 12-min Cooper test. Sun: measurements, photos, review."),
        ("Tue/Thu", "Zone 2 cardio (talk test, ~60-70% HRmax): 30 min weeks 1-4, 35-40 min from week 5. Incline treadmill, bike, rower or elliptical."),
        ("Sat", "45-60 min of an activity you enjoy."),
        ("Sun", "20 min mobility (hips, thoracic spine, shoulders, wrists)."),
        ("Steps", "8,000/day weeks 0-4, 10,000/day from week 5 (15-min walks after lunch and dinner)."),
    ], merge_to=5)
    ws.cell(row=r, column=1, value="Wrist and grip strategy").font = BOLD
    r = write_table(ws, r + 1, ["#", "Rule"], [
        (1, "Neutral-grip dumbbells for pressing in weeks 1-4; barbells only when wrists feel stable, knuckles stacked over the forearm."),
        (2, "Wrist warm-up every session; wrist wraps only for heavy pressing if there is discomfort."),
        (3, "Two grip finishers per week (farmer carry, dead hang)."),
        (4, "Lifting straps from week 8 on RDL, deadlift and heavy rows so grip fatigue never limits leg and back development."),
    ], merge_to=5)
    paragraph(ws, r, "Cycle 2 (Weeks 14-26): four-day upper/lower split. Mon Lower A, Tue Upper A, Wed Zone 2 40 min, Thu Lower B, Fri Upper B, Sat intervals (8 x 1 min hard / 2 min easy) or sport, Sun mobility.", 1, 5, bold=True)
    write_table(ws, r + 2, ["Session", "Exercises (sets x reps)"], [
        ("Lower A", "Back squat 4 x 5-8; Romanian deadlift 3 x 8-10; leg press or split squat 3 x 10-12; calf raise 3 x 12-15; hanging knee raise 3 x 10"),
        ("Upper A", "Bench press 4 x 5-8; chin-up or lat pulldown 4 x 6-10; dumbbell shoulder press 3 x 8-12; cable row 3 x 10-12; face pull 3 x 15; farmer carry 3 x 40 m"),
        ("Lower B", "Trap-bar or conventional deadlift 4 x 5 (straps allowed); front squat or hack squat 3 x 8-10; hip thrust 3 x 10-12; leg curl 3 x 12; plank 3 x 60 s"),
        ("Upper B", "Incline dumbbell press 4 x 8-10; pull-up or pulldown 4 x 8-10; dips or close-grip push-up 3 x 8-12; one-arm dumbbell row 3 x 10-12; lateral raise 3 x 12-15; dead hang 3 x max"),
        ("Progression", "Weeks 14-21 double progression; week 22 deload; weeks 23-25 main lifts 4 x 5-6; week 26 final benchmarks (same tests as week 13)."),
    ], merge_to=5)


def build_training_log(wb):
    ws = wb.create_sheet("TrainingLog")
    headers = ["Date", "Weekday", "Week", "Planned session", "Done (1/0)", "Duration (min)",
               "Lift 1 (exercise: kg x reps)", "Lift 2", "Lift 3", "RPE (1-10)", "Notes"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, 1, len(headers))
    ws.row_dimensions[1].height = 32
    for i in range(N_DAYS):
        d = START + timedelta(days=i)
        r = 2 + i
        ws.cell(row=r, column=1, value=d).number_format = DATE_FMT
        ws.cell(row=r, column=2, value=d.strftime("%a"))
        ws.cell(row=r, column=3, value=week_of(d))
        ws.cell(row=r, column=4, value=planned_training(d))
        for col in range(5, 12):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
    last = N_DAYS + 1
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{last}")
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=["1"], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=["0"], fill=RED_FILL))
    widths = [14, 8, 6, 60, 9, 9, 26, 26, 26, 9, 36]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "E2"


def build_skill_plan(wb):
    ws = wb.create_sheet("SkillPlan")
    set_widths(ws, {"A": 6, "B": 16, "C": 24, "D": 62, "E": 44, "F": 40, "G": 9, "H": 9, "I": 30})
    title(ws, "Skill block: GIS and remote sensing", "60 min Mon-Fri (20:45-21:45), optional 90 min Saturday project session, one commit pushed every Friday. One study area for all modules.")
    headers = ["Week", "Dates", "Module", "Focus this week", "Resources (free unless noted)", "Friday deliverable", "Done (Y/N)", "Hours logged", "Notes / link"]
    rows = []
    for (w, module, focus, res, deliverable) in SKILL_PLAN:
        s, e = week_dates(w)
        rows.append((w, f"{s:%d %b} - {e:%d %b}", module, focus, res, deliverable, None, None, None))
    hr = 4
    r = write_table(ws, hr, headers, rows)
    first, last = hr + 1, hr + len(rows)
    for rr in range(first, last + 1):
        for col in (7, 8, 9):
            ws.cell(row=rr, column=col).fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G{first}:G{last}")
    ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=['"Y"'], fill=GREEN_FILL))
    paragraph(ws, r, "Certification: primary Esri ArcGIS Pro Associate 2025 (75 questions, 90 min; aimed at 2-4 years of experience) or ArcGIS Pro Professional 2025 if you have more than 4 years. Stretch option matching the new skills: ArcGIS API for Python Associate 2026. Buy the voucher and book the date in week 20. Optional paid cohort: Spatial Thoughts live 'Python Foundation for Spatial Analysis', 21-22 and 28-29 Oct 2026 (aligns with weeks 3-4). NASA ARSET trainings are free: join any GEE or land-cover session scheduled Oct-Mar.", 1, 6)
    r += 2
    ws.cell(row=r, column=4, value="Key free resources").font = BOLD
    write_table(ws, r + 1, ["Resource", "URL"], [
        ("Spatial Thoughts OpenCourseWare (Python Foundation, End-to-End GEE, Mapping and Data Visualization with Python, Cloud Native Remote Sensing with Python)", "https://courses.spatialthoughts.com/"),
        ("Cloud-Based Remote Sensing with Google Earth Engine: Fundamentals and Applications (open access)", "https://link.springer.com/book/10.1007/978-3-031-26588-4"),
        ("NASA ARSET - Using Google Earth Engine for Land Monitoring Applications", "https://www.earthdata.nasa.gov/learn/trainings/using-google-earth-engine-land-monitoring-applications"),
        ("NASA ARSET training catalogue", "https://www.earthdata.nasa.gov/learn/trainings"),
        ("Introduction to PostGIS workshop", "https://postgis.net/workshops/postgis-intro/"),
        ("GeoPandas documentation", "https://geopandas.org/"),
        ("Rasterio documentation", "https://rasterio.readthedocs.io/"),
        ("geemap documentation", "https://geemap.org/"),
        ("Esri Technical Certification (exam list, learning plans)", "https://www.esri.com/training/certification/"),
    ], start_col=4)
    ws.freeze_panes = f"C{first}"


def build_meals(wb, t):
    ws = wb.create_sheet("Meals")
    set_widths(ws, {"A": 18, "B": 30, "C": 84, "D": 12, "E": 12})
    title(ws, "Meal framework", "Numbers come from the Targets sheet. Protein first on every plate; oil measured with a spoon; no liquid calories.")
    r = 4
    ws.cell(row=r, column=1, value="Live targets").font = BOLD
    live = [
        ("Calories (kcal)", f"=Targets!{t['kcal']}", "0", "Weeks 0-13 deficit. Week 13 checkpoint: waist < 87 cm -> set deficit to 0 on Targets. Week 22: deficit 0 for 7 days."),
        ("Protein (g)", f"=Targets!{t['protein']}", "0", "130-150 g; spread over 4-5 meals"),
        ("Fat (g)", f"=Targets!{t['fat']}", "0", ""),
        ("Carbohydrate (g)", f"=Targets!{t['carbs']}", "0", "Breakfast, pre-gym, dinner"),
        ("Water, training day (L)", f"=Targets!{t['water']}", "0.0", "500 ml on waking, 500 ml pre-gym, 500-750 ml during training; rest days 0.5 L less"),
        ("Fibre (g)", 30, "0", "Vegetables, legumes, oats, fruit"),
        ("Sugary drinks / juice", 0, "0", "Fastest belly-fat lever"),
        ("Alcohol", "None or social only", "@", "Decide before Day 1"),
        ("Sweets", "0 in Oct; 2 planned/week from Nov", "@", "Planned treats are not misses"),
        ("Expected rate", "0.3-0.5 kg/week; waist ~1 cm / 2 weeks", "@", "Slower on the scale if muscle is gained"),
    ]
    top = r + 1
    r = write_table(ws, top, ["Target", "Value", "Note"], [(a, b, d) for a, b, _f, d in live])
    for i, (_a, _b, fmt, _d) in enumerate(live, start=1):
        ws.cell(row=top + i, column=2).number_format = fmt
    ws.cell(row=r, column=1, value="Daily template: base 2,150 kcal / 145 g protein; the optional 21:30 snack adds 150 kcal / 15 g. To land near the calorie target, skip the optional snack and halve the lunch starch.").font = BOLD
    meals = [
        ("06:45", "Breakfast", "3 eggs or 250 g strained yogurt; 60 g oats or 2 slices wholegrain bread; 1 fruit; tea/coffee without sugar", 500, 35),
        ("10:30", "Snack", "200 g yogurt, or 30 g nuts + fruit, or 2 boiled eggs", 200, 15),
        ("12:30", "Lunch (packed)", "150 g cooked chicken/fish/lean meat or 200 g legumes; 1 cup cooked rice/bulgur or 2 flatbreads; large salad or cooked vegetables; 1 tbsp olive oil", 600, 40),
        ("16:30", "Pre-gym", "Banana + 3 dates, or bread with cheese/peanut butter; 500 ml water", 250, 10),
        ("18:00-19:15", "Gym", "500-750 ml water", 0, 0),
        ("19:45", "Dinner", "150-200 g protein source; 2 cups vegetables; one fist of potatoes or rice; 1 tbsp oil", 600, 45),
        ("21:30 (optional)", "If hungry", "Cottage cheese, yogurt or a protein shake; kitchen closed at 22:00", 150, 15),
    ]
    top = r + 1
    first, last = top + 1, top + len(meals)
    totals = [
        ("", "Total, base (without optional snack)", "", f"=SUM(D{first}:D{last - 1})", f"=SUM(E{first}:E{last - 1})"),
        ("", "Total with optional snack", "", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})"),
    ]
    r = write_table(ws, top, ["Time", "Meal", "Example", "kcal", "Protein (g)"], meals + totals)
    for rr in (last + 1, last + 2):
        for col in (2, 4, 5):
            ws.cell(row=rr, column=col).font = BOLD
    ws.cell(row=r, column=1, value="Rules").font = BOLD
    r = write_table(ws, r + 1, ["#", "Rule"], [
        (1, "Protein first on every plate; vegetables at lunch and dinner; oil measured with a spoon."),
        (2, "Eat out at most twice a week; order protein plus vegetables first; the rest is 80/20."),
        (3, "Sunday prep (60 min): cook 1 kg chicken or legumes, 4 cups rice, chop vegetables for 4 days, boil 8 eggs, portion 5 lunches."),
        (4, "Adjust every 2 weeks, not every day: if weight and waist have not moved in 14 days, remove 150 kcal (the dinner starch) or add 1,500 steps. One change at a time."),
        (5, "Week 13 checkpoint: waist below 87 cm -> maintenance (deficit 0 on Targets) and train for muscle; otherwise continue the deficit."),
        (6, "Week 22: 7-day diet break at maintenance, deliberately, together with the training deload."),
        (7, "Sleep 7-8 h. Short sleep raises appetite and stores fat around the waist."),
    ], merge_to=3)
    ws.cell(row=r, column=1, value="Sunday prep checklist").font = BOLD
    r = write_table(ws, r + 1, ["Item", "Done"], [
        ("1 kg chicken / fish / legumes cooked", None), ("4 cups rice or bulgur cooked", None), ("Vegetables chopped for 4 days", None),
        ("8 eggs boiled", None), ("5 lunches portioned", None), ("Yogurt, fruit, nuts for snacks bought", None), ("Water bottle (1 L) ready for work and gym", None),
    ], start_col=2)


def build_monthly_review(wb):
    ws = wb.create_sheet("MonthlyReview")
    set_widths(ws, {"A": 20, "B": 14, "C": 16, "D": 12, "E": 30, "F": 30, "G": 34, "H": 30, "I": 30, "J": 14})
    title(ws, "Monthly review (30 min) and cycle reviews (60 min)", "Month completion % is computed from DailyLog. Read the whole plan before the Week 13 and Week 26 reviews; do not revise from memory.")
    months = [
        ("M1", date(2026, 10, 1), date(2026, 11, 1)),
        ("M2", date(2026, 11, 2), date(2026, 11, 29)),
        ("M3 (Cycle 1 review)", date(2026, 11, 30), date(2027, 1, 3)),
        ("M4", date(2027, 1, 4), date(2027, 1, 31)),
        ("M5", date(2027, 2, 1), date(2027, 2, 28)),
        ("M6 (final review)", date(2027, 3, 1), date(2027, 4, 4)),
    ]
    headers = ["Month", "From", "Review date (Sunday)", "Completion %", "What worked", "What broke", "Which metrics moved (waist, weight, lifts, skill)",
               "One change for next month", "Rules to simplify or drop", "Measurements + photos done (Y/N)"]
    hr = 4
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=j, value=h)
    style_header(ws, hr, 1, len(headers))
    ws.row_dimensions[hr].height = 40
    dates = f"DailyLog!$B${DL_FIRST}:$B${DL_LAST}"
    for i, (name, s, e) in enumerate(months, start=1):
        r = hr + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=s).number_format = DATE_FMT
        ws.cell(row=r, column=3, value=e).number_format = DATE_FMT
        ws.cell(row=r, column=4, value=(
            f'=IFERROR(SUMIFS(DailyLog!$O${DL_FIRST}:$O${DL_LAST},{dates},">="&B{r},{dates},"<="&C{r})'
            f'/SUMIFS(DailyLog!$R${DL_FIRST}:$R${DL_LAST},{dates},">="&B{r},{dates},"<="&C{r}),"")')).number_format = "0%"
        for col in range(5, 11):
            ws.cell(row=r, column=col).fill = INPUT_FILL
            ws.cell(row=r, column=col).alignment = WRAP
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
        ws.row_dimensions[r].height = 70
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J{hr + 1}:J{hr + len(months)}")
    r = hr + len(months) + 2
    ws.cell(row=r, column=1, value="Cycle review questions (Week 13 and Week 26)").font = BOLD
    qs = [
        "Do I have the time and energy for each rule as written? If not, what is the baseline version?",
        "Am I getting results that match the effort? If effort is high and results are flat, change the approach, not the goal.",
        "Does each rule still fit my life and values? Keep, simplify or drop.",
        "What evidence of progress do I have (waist, lifts, deliverables, books)? Compare to Day 1, not to the ideal.",
        "Nutrition checkpoint: is the waist below 87 cm? If yes, set the deficit to 0 and train for muscle.",
        "Skill: what is public in the portfolio? What is the next deliverable that would matter to an employer or client?",
        "What is the one bold move for the next 13 weeks?",
    ]
    for i, q in enumerate(qs, start=1):
        paragraph(ws, r + i, f"{i}. {q}", 1, 10)


def print_setup(wb):
    """Landscape, fit to one page wide, header row repeated on the long logs."""
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
    wb["DailyLog"].print_title_rows = "1:1"
    wb["TrainingLog"].print_title_rows = "1:1"


def main():
    wb = Workbook()
    build_overview(wb)
    targets_ref = build_targets(wb)
    build_daily_log(wb)
    build_weekly_review(wb)
    build_benchmarks(wb)
    build_training_plan(wb)
    build_training_log(wb)
    build_skill_plan(wb)
    build_meals(wb, targets_ref)
    build_monthly_review(wb)
    print_setup(wb)
    wb.save(OUT)
    print(f"Wrote {OUT} ({N_DAYS} days, weeks 0-{N_WEEKS}, {START} to {END})")


if __name__ == "__main__":
    main()
