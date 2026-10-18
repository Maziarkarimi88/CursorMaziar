#!/usr/bin/env python3
"""Parse the WBK/AF-2 master contract workbook into imported_schemes.json."""

from __future__ import annotations

import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

_MASTER_CANDIDATES = [
    Path("/workspace/Copy of WBK-AF2_Master file contract status awarded and remaining contracts-rev.xlsx"),
    Path("/workspace/2- WBK-AF2_Master file contract status awarded and remaining contracts.xlsx"),
]
MASTER = next(p for p in _MASTER_CANDIDATES if p.exists())
OUT = Path(__file__).resolve().parent / "imported_schemes.json"

PROJECT = "Emergency Food Security Project (EFSP) – WBK/AF-2"
DONOR = "World Bank (WBK/AF-2)"

REGION_MAP = {
    "west/herat": "West / Herat",
    "west/herat region": "West / Herat",
    "east/jalalabad": "East / Jalalabad",
    "central/kabul": "Central / Kabul",
    "central /kabul": "Central / Kabul",
    "central highland/bamyan": "Central Highland / Bamyan",
    "central highlan/ bamyan": "Central Highland / Bamyan",
    "central highland/bamayn": "Central Highland / Bamyan",
    "southeast/gardez": "Southeast / Gardez",
    "south-east/gardez": "Southeast / Gardez",
    "southeastern/gardez": "Southeast / Gardez",
    "south/kandahar": "South / Kandahar",
    "southern/kandahar": "South / Kandahar",
    "north/mazar": "North / Mazar",
    "northern/mazar": "North / Mazar",
    "northeast/kunduz": "Northeast / Kunduz",
    "north-east/kunduz": "Northeast / Kunduz",
    "kabul": "Central / Kabul",
    "jalalabad": "East / Jalalabad",
    "herat": "West / Herat",
    "kandahar": "South / Kandahar",
    "kunduz": "Northeast / Kunduz",
    "kundoz": "Northeast / Kunduz",
    "mazar": "North / Mazar",
    "bamiyan": "Central Highland / Bamyan",
    "gardiz": "Southeast / Gardez",
    "gardez": "Southeast / Gardez",
    "eastern": "East / Jalalabad",
}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v
    s = str(v).replace("\n", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def parse_date(ws, row, col, cache):
    key = (ws.title, row, col)
    if key in cache:
        return cache[key]
    v = ws.cell(row, col).value
    result = None
    if isinstance(v, datetime):
        result = v
    elif isinstance(v, date):
        result = datetime(v.year, v.month, v.day)
    elif isinstance(v, (int, float)) and v > 20000:
        try:
            from openpyxl.utils.datetime import from_excel

            result = from_excel(v)
        except Exception:
            result = None
    elif isinstance(v, str):
        s = v.strip()
        m = re.match(r"^=([A-Z]+)(\d+)\+(\d+)$", s.replace(" ", "").upper())
        if m:
            src = parse_date(ws, int(m.group(2)), column_index_from_string(m.group(1)), cache)
            if src:
                result = src + timedelta(days=int(m.group(3)))
        else:
            s2 = s.replace("Oc-", "Oct-").replace("oc-", "Oct-")
            for fmt in ("%Y-%m-%d", "%d-%b-%y", "%d-%b-%Y", "%m/%d/%Y", "%d/%m/%Y", "%b, %Y", "%b. %Y"):
                try:
                    result = datetime.strptime(s2.replace(",", ""), fmt) if fmt.startswith("%b") else datetime.strptime(s2, fmt)
                    break
                except ValueError:
                    continue
            if result is None:
                m2 = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
                if m2:
                    result = datetime(int(m2.group(3)), int(m2.group(1)), int(m2.group(2)))
    cache[key] = result
    return result


def norm_region(v):
    s = clean(v)
    if not s or isinstance(s, datetime):
        return ""
    return REGION_MAP.get(s.lower().strip(), s)


def norm_status(v):
    s = clean(v)
    if not s:
        return ""
    key = s.lower().strip(" .")
    mapping = {
        "under design": "Under Design",
        "under wri review (design review)": "Under Review of WRI",
        "design under wri review": "Under Review of WRI",
        "under lto review": "Under LTO Review",
        "pr to be raised": "PR to be Raised",
        "itb launched": "ITB Launched",
        "under te": "Under Technical Evaluation",
        "te completed": "Under Technical Evaluation",
        "retendered (reitb)": "Retendered (Re-ITB)",
        "retendered (re-itb)": "Retendered (Re-ITB)",
        "retentered (re-itb)": "Retendered (Re-ITB)",
    }
    return mapping.get(key, s)


def to_progress(v):
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    if n <= 1:
        n = n * 100
    return round(n, 1)


def to_cost(v):
    if v is None or isinstance(v, str) and v.startswith("="):
        return None
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def iso(d):
    if not d:
        return None
    if isinstance(d, datetime):
        return d.strftime("%Y-%m-%d")
    return None


def months_between(start, end):
    if not start or not end:
        return None
    days = (end - start).days
    if days <= 0:
        return None
    return round(days / 30.4, 1)


def skip_row(name):
    s = clean(name).lower()
    return (not s) or s.startswith("total") or s.startswith("subtotal") or s.startswith("sub total") or s.startswith("grand")


def carry(pkg, row, fields):
    for f in fields:
        if row.get(f):
            pkg[f] = row[f]
        elif pkg.get(f) and not row.get(f):
            row[f] = pkg[f]
    return pkg


def extract_standard(ws, *, start_row, name_col, award_group, family, intervention, awarded):
    cache = {}
    rows = []
    pkg = {}
    for r in range(start_row, (ws.max_row or 0) + 1):
        name = clean(ws.cell(r, name_col).value)
        if skip_row(name):
            continue
        if awarded:
            package = clean(ws.cell(r, 2).value)
            region = norm_region(ws.cell(r, 4).value)
            province = clean(ws.cell(r, 5).value)
            district = clean(ws.cell(r, 6).value)
            village = clean(ws.cell(r, 7).value)
            signed = parse_date(ws, r, 8, cache)
            end = parse_date(ws, r, 9, cache)
            if family.startswith("Watershed"):
                cost = to_cost(ws.cell(r, 10).value)
                number = clean(ws.cell(r, 11).value)
                contractor = clean(ws.cell(r, 12).value)
                progress = to_progress(ws.cell(r, 13).value)
                remarks = clean(ws.cell(r, 16).value)
                actual = None
            else:
                number = clean(ws.cell(r, 10).value)
                cost = to_cost(ws.cell(r, 11).value)
                contractor = clean(ws.cell(r, 12).value)
                progress = to_progress(ws.cell(r, 13).value)
                actual_raw = clean(ws.cell(r, 15).value)
                actual = parse_date(ws, r, 15, cache)
                remarks = clean(ws.cell(r, 16).value)
                if actual_raw and not actual:
                    remarks = (remarks + " | Actual completion: " + actual_raw).strip(" |")
            rec = {
                "package_name": package,
                "scheme_name": name,
                "region": region,
                "province": province,
                "district": district,
                "village": village,
                "awarded_date": iso(signed),
                "contract_end_date": iso(end),
                "actual_completion_date": iso(actual) if actual else None,
                "contract_number": number,
                "contract_cost_usd": cost,
                "contractor_name": contractor,
                "physical_progress_pct": progress,
                "remarks": remarks,
                "contract_duration_months": months_between(signed, end),
            }
            pkg = carry(pkg, rec, [
                "package_name", "region", "awarded_date", "contract_end_date",
                "contract_number", "contractor_name", "remarks",
                "contract_duration_months",
            ])
            status = "Awarded"
            remarks_l = (rec["remarks"] or "").lower()
            completed_note = remarks_l.startswith("completed in") or remarks_l.startswith("completed.")
            if "cancel" in remarks_l:
                status = "Cancelled"
            elif (progress or 0) >= 100 or completed_note:
                status = "Completed / Handover"
            elif (progress or 0) > 0:
                status = "Under Construction"
            rec["reported_status"] = status
            rec["award_group"] = award_group or "Awarded"
            rec["intervention_family"] = family
            rec["intervention_type"] = intervention
            rec["source_sheet"] = ws.title
            rows.append(rec)
        else:
            package = clean(ws.cell(r, 2).value)
            region = norm_region(ws.cell(r, 4).value)
            province = clean(ws.cell(r, 5).value)
            district = clean(ws.cell(r, 6).value)
            village = clean(ws.cell(r, 7).value)
            status = norm_status(ws.cell(r, 8).value)
            exp_award = parse_date(ws, r, 9, cache)
            exp_end = parse_date(ws, r, 10, cache)
            cost = to_cost(ws.cell(r, 11).value)
            remarks = clean(ws.cell(r, 12).value)
            rec = {
                "package_name": package,
                "scheme_name": name,
                "region": region,
                "province": province,
                "district": district,
                "village": village,
                "reported_status": status,
                "expected_award_date": iso(exp_award),
                "contract_end_date": iso(exp_end),
                "contract_cost_usd": cost,
                "remarks": remarks,
                "contract_duration_months": months_between(exp_award, exp_end),
            }
            pkg = carry(pkg, rec, [
                "package_name", "region", "reported_status", "expected_award_date",
                "contract_end_date", "contract_cost_usd", "remarks",
                "contract_duration_months",
            ])
            rec["award_group"] = award_group or "Not awarded"
            rec["intervention_family"] = family
            rec["intervention_type"] = intervention
            rec["source_sheet"] = ws.title
            rows.append(rec)
    return rows


def extract_drip(ws):
    cache = {}
    rows = []
    pkg = {}
    for r in range(9, (ws.max_row or 0) + 1):
        name = clean(ws.cell(r, 6).value)
        if skip_row(name) or name.lower() in {"sub total", "name of the sub-project"}:
            continue
        package = clean(ws.cell(r, 2).value)
        code = clean(ws.cell(r, 5).value)
        region = norm_region(ws.cell(r, 7).value)
        province = clean(ws.cell(r, 8).value)
        district = clean(ws.cell(r, 9).value)
        village = clean(ws.cell(r, 10).value)
        farmer = clean(ws.cell(r, 11).value)
        cost = to_cost(ws.cell(r, 14).value)
        status = norm_status(ws.cell(r, 15).value)
        exp_award = parse_date(ws, r, 16, cache)
        exp_end = parse_date(ws, r, 17, cache)
        remarks = clean(ws.cell(r, 18).value)
        rec = {
            "package_name": package,
            "scheme_name": name,
            "scheme_code": code,
            "region": region,
            "province": province,
            "district": district,
            "village": village,
            "contractor_name": farmer,
            "contract_cost_usd": cost,
            "reported_status": status,
            "expected_award_date": iso(exp_award),
            "contract_end_date": iso(exp_end),
            "remarks": remarks,
            "contract_duration_months": months_between(exp_award, exp_end),
        }
        pkg = carry(pkg, rec, [
            "package_name", "region", "reported_status", "expected_award_date",
            "contract_end_date", "remarks", "contract_duration_months",
        ])
        rec["award_group"] = "Not awarded"
        rec["intervention_family"] = "Drip irrigation"
        rec["intervention_type"] = "Drip Irrigation"
        rec["source_sheet"] = ws.title
        if farmer:
            rec["remarks"] = ((rec.get("remarks") or "") + (f" | Farmer: {farmer}" if farmer else "")).strip(" |")
        rows.append(rec)
    return rows


def extract_lll(ws):
    cache = {}
    rows = []
    pkg = {}
    for r in range(19, (ws.max_row or 0) + 1):
        name = clean(ws.cell(r, 3).value)
        if skip_row(name) or name.startswith("S."):
            continue
        region = norm_region(ws.cell(r, 4).value)
        province = clean(ws.cell(r, 5).value)
        district = clean(ws.cell(r, 6).value)
        units = ws.cell(r, 7).value
        award = parse_date(ws, r, 8, cache)
        supply = parse_date(ws, r, 9, cache)
        dist = parse_date(ws, r, 10, cache)
        remarks = clean(ws.cell(r, 11).value)
        rec = {
            "package_name": "EFSP-AF2 LLL supply and delivery (50 units)",
            "scheme_name": name,
            "region": region,
            "province": province,
            "district": district,
            "expected_award_date": iso(award),
            "contract_end_date": iso(dist or supply),
            "remarks": remarks,
            "physical_progress_pct": None,
        }
        pkg = carry(pkg, rec, ["expected_award_date", "contract_end_date", "remarks"])
        rec["award_group"] = "Not awarded"
        rec["intervention_family"] = "LLL supply"
        rec["intervention_type"] = "Laser Land Leveler (LLL)"
        rec["source_sheet"] = ws.title
        rec["reported_status"] = "Under Awarding"
        rec["village"] = f"{units} unit(s)" if units else ""
        rec["contract_duration_months"] = months_between(award, dist or supply)
        rows.append(rec)
    return rows


def extract_agro(ws):
    cache = {}
    rows = []
    pkg = {}
    for r in range(21, (ws.max_row or 0) + 1):
        region = norm_region(ws.cell(r, 2).value)
        province = clean(ws.cell(r, 3).value)
        district = clean(ws.cell(r, 4).value)
        if skip_row(province) or not province:
            continue
        saplings = ws.cell(r, 5).value
        seeds = ws.cell(r, 6).value
        remarks = clean(ws.cell(r, 9).value)
        rec = {
            "package_name": "Agroforestry saplings and seeds 2026/27",
            "scheme_name": f"Agroforestry — {province} / {district}",
            "region": region,
            "province": province,
            "district": district,
            "village": f"Saplings: {saplings or 0}; Seeds kg: {seeds or 0}",
            "remarks": remarks,
            "expected_award_date": "2026-04-30",
            "contract_end_date": "2027-04-01",
        }
        pkg = carry(pkg, rec, ["remarks", "region"])
        rec["award_group"] = "Not awarded"
        rec["intervention_family"] = "Agroforestry"
        rec["intervention_type"] = "Agroforestry saplings / seeds"
        rec["source_sheet"] = ws.title
        rec["reported_status"] = "Under Awarding"
        rec["contract_duration_months"] = 12
        rows.append(rec)
    return rows


MAJOR_FROM_FAMILY = {
    "On/off-farm irrigation": "Irrigation Scheme",
    "Watershed / Check dam": "Watershed Management",
    "Drip irrigation": "Drip Irrigation",
    "LLL supply": "Laser Land Leveler (LLL)",
    "Agroforestry": "Agroforestry saplings / seeds",
}

MAJOR_CODE = {
    "Irrigation Scheme": "IS",
    "Watershed Management": "WSM",
    "Drip Irrigation": "DIS",
    "Laser Land Leveler (LLL)": "LLL",
    "Agroforestry saplings / seeds": "AGF",
}


def infer_subtype(name: str, major: str) -> str:
    if major == "Laser Land Leveler (LLL)":
        return "LLL unit"
    if major == "Agroforestry saplings / seeds":
        return "Saplings and seeds"
    if major == "Drip Irrigation":
        return "Drip Irrigation Plot"
    n = (name or "").lower()
    if "protection wall" in n or "retaining wall" in n:
        return "Protection Wall"
    if "karez" in n:
        return "Karez Rehabilitation"
    if "pond" in n and "check" in n:
        return "Check Dam and Water Pond"
    if "boulder" in n:
        return "Boulder Check Dam"
    if "check dam" in n or "checkdam" in n:
        return "Check Dam"
    if "drip" in n:
        return "Drip Irrigation Plot"
    if "canal" in n or "irrigation" in n:
        return "Irrigation Canal"
    if major == "Watershed Management":
        return "Check Dam"
    if major == "Irrigation Scheme":
        return "Irrigation Canal"
    return "Other"


def make_package_id(package_name: str, major: str) -> str:
    code = MAJOR_CODE.get(major, "X")
    text = package_name or ""
    m = re.search(r"Package\s*#\s*(\d+)\s*([A-Za-z])?", text, re.I)
    if m:
        num = m.group(1).zfill(2)
        letter = (m.group(2) or "").upper()
        return f"AF2-{code}-{num}{letter}"
    if code == "LLL":
        return "AF2-LLL-ALL"
    if code == "AGF":
        return "AF2-AGF-ALL"
    slug = re.sub(r"[^A-Za-z0-9]+", "", text)[:12] or "MISC"
    return f"AF2-{code}-{slug}"


def assign_ids(rows):
    for rec in rows:
        if rec.get("award_group") in ("Remaining", "remaining", None, ""):
            rec["award_group"] = "Not awarded"
        major = MAJOR_FROM_FAMILY.get(rec.get("intervention_family", ""), rec.get("intervention_type") or "")
        rec["intervention_major"] = major
        rec["intervention_type"] = infer_subtype(rec.get("scheme_name", ""), major)
        rec["package_id"] = make_package_id(rec.get("package_name", ""), major)
        rec["project_name"] = PROJECT
        rec["donor"] = DONOR
        rec.setdefault("contractor_name", "")
        rec.setdefault("contract_number", "")
        rec.setdefault("village", "")
        rec.setdefault("physical_progress_pct", None)
        rec.setdefault("expected_award_date", None)
        rec.setdefault("awarded_date", None)
        rec.setdefault("actual_completion_date", None)
        rec.setdefault("status_dates", {})
        if rec.get("awarded_date"):
            rec["status_dates"]["Awarded"] = rec["awarded_date"]
        if rec.get("reported_status") == "Completed / Handover":
            done = rec.get("actual_completion_date") or rec.get("contract_end_date")
            if done:
                rec["status_dates"]["Completed / Handover"] = done
            if rec.get("awarded_date"):
                rec["status_dates"]["Under Construction"] = rec["awarded_date"]
        elif rec.get("reported_status") == "Under Construction" and rec.get("awarded_date"):
            rec["status_dates"]["Under Construction"] = rec["awarded_date"]

    rows.sort(key=lambda r: (
        r.get("intervention_major") or "",
        r.get("package_id") or "",
        r.get("scheme_name") or "",
    ))

    counts = {}
    pkg_sizes = {}
    for rec in rows:
        pid = rec["package_id"]
        pkg_sizes[pid] = pkg_sizes.get(pid, 0) + 1
    pkg_seen = {}
    for rec in rows:
        code = MAJOR_CODE.get(rec["intervention_major"], "X")
        grp = "A" if rec["award_group"] == "Awarded" else "N"
        key = f"{code}-{grp}"
        counts[key] = counts.get(key, 0) + 1
        rec["scheme_id"] = f"AF2-{key}-{counts[key]:03d}"
        pid = rec["package_id"]
        pkg_seen[pid] = pkg_seen.get(pid, 0) + 1
        rec["subproject_no"] = f"{pkg_seen[pid]} of {pkg_sizes[pid]}"
        rec["subproject_count"] = pkg_sizes[pid]
    return rows


def main():
    wb = load_workbook(MASTER, data_only=False)
    rows = []
    rows += extract_standard(
        wb["1.Awarded contracts_On&off-farm"],
        start_row=8, name_col=3, award_group="Awarded",
        family="On/off-farm irrigation", intervention="Irrigation Canal", awarded=True,
    )
    rows += extract_standard(
        wb["2.Remain contracts_On&off-farm"],
        start_row=7, name_col=3, award_group="Not awarded",
        family="On/off-farm irrigation", intervention="Irrigation Canal", awarded=False,
    )
    rows += extract_standard(
        wb["3.WSM awarded contracts"],
        start_row=8, name_col=3, award_group="Awarded",
        family="Watershed / Check dam", intervention="Check Dam", awarded=True,
    )
    rows += extract_standard(
        wb["4.WSM remaining contracts"],
        start_row=8, name_col=3, award_group="Not awarded",
        family="Watershed / Check dam", intervention="Check Dam", awarded=False,
    )
    rows += extract_drip(wb["5. Drip Irrigation systems"])
    rows += extract_lll(wb["6. Supply and delivery of LLL"])
    rows += extract_agro(wb["7. Agroforestry Saplings-seeds"])
    rows = assign_ids(rows)
    OUT.write_text(json.dumps(rows, indent=2, ensure_ascii=False))
    from collections import Counter
    print(f"Wrote {len(rows)} schemes to {OUT}")
    print("by major", Counter(r["intervention_major"] for r in rows))
    print("by subtype", Counter(r["intervention_type"] for r in rows))
    print("by group", Counter(r["award_group"] for r in rows))
    print("by status", Counter(r.get("reported_status") or "(blank)" for r in rows).most_common(15))
    multi = Counter(r["package_id"] for r in rows)
    print("multi-scheme packages", sum(1 for n in multi.values() if n > 1), "largest", multi.most_common(5))
    pkg_key = {}
    for r in rows:
        pid = r["package_id"]
        if pid not in pkg_key:
            pkg_key[pid] = (r["intervention_major"], r["award_group"])
    print("packages by major × group")
    print(Counter(pkg_key.values()))
    print(f"source {MASTER.name}")


if __name__ == "__main__":
    main()
